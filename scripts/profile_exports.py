"""Read-only source profile for the six exported product modules."""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ID_PATTERN = re.compile(r"/e-commerce/detail/(\d+)")


def profile(path: Path) -> dict:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    headers = [str(cell.value or "") for cell in next(sheet.iter_rows(max_row=1))]
    id_columns = [index for index, name in enumerate(headers) if "商品ID" in name]
    detail_columns = [
        index for index, name in enumerate(headers)
        if "FastMoss" in name and "商品详情页" in name
    ]
    image_columns = [
        index for index, name in enumerate(headers)
        if "商品图片" in name or "商品封面" in name
    ]
    missing = Counter()
    ids = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        product_id = str(row[id_columns[0]]) if id_columns and row[id_columns[0]] else None
        if not product_id and detail_columns:
            match = ID_PATTERN.search(str(row[detail_columns[0]] or ""))
            product_id = match.group(1) if match else None
        if product_id:
            ids.append(product_id)
        else:
            missing["product_id"] += 1
        if not image_columns or not row[image_columns[0]]:
            missing["image_url"] += 1
    return {
        "file": path.name,
        "rows": sheet.max_row - 1,
        "id_columns": [headers[i] for i in id_columns],
        "detail_columns": [headers[i] for i in detail_columns],
        "image_columns": [headers[i] for i in image_columns],
        "unique_ids": len(set(ids)),
        "duplicate_ids": len(ids) - len(set(ids)),
        "missing": dict(missing),
        "first_row": {
            name: value for name, value in zip(headers, next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
            if name and (name in [headers[i] for i in id_columns + detail_columns + image_columns] or name in {"商品", "商品名称", "商品标题"})
        },
    }


if __name__ == "__main__":
    for raw in sys.argv[1:]:
        print(json.dumps(profile(Path(raw)), ensure_ascii=True, default=str))
