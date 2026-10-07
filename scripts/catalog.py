"""Import FastMoss's own product exports and retain verified product images.

The source XLSX files are never modified. Every exported row and image result is
tracked in SQLite so an interrupted run can resume without silently skipping work.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sqlite3
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PRODUCT_ID = re.compile(r"/e-commerce/detail/(\d+)")
DEFAULT_ROOT = Path(os.environ.get("FASTMOSS_DATA_ROOT", str(Path(__file__).resolve().parents[1] / "data")))


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def connect(root: Path) -> sqlite3.Connection:
    root.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(root / "catalog.sqlite3", timeout=30)
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS source_exports (
            source_file TEXT PRIMARY KEY,
            sha256 TEXT NOT NULL,
            region TEXT NOT NULL,
            module TEXT NOT NULL,
            snapshot_date TEXT NOT NULL,
            first_rank INTEGER NOT NULL,
            last_rank INTEGER NOT NULL,
            row_count INTEGER NOT NULL,
            imported_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS products (
            product_id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            status TEXT,
            shop_name TEXT,
            region TEXT,
            category TEXT,
            price TEXT,
            commission_rate TEXT,
            day7_sold REAL,
            day7_gmv REAL,
            total_sold REAL,
            total_gmv REAL,
            author_count REAL,
            author_order_rate TEXT,
            video_count REAL,
            live_count REAL,
            image_url TEXT NOT NULL,
            fastmoss_detail_url TEXT NOT NULL,
            tiktok_detail_url TEXT,
            shop_detail_url TEXT,
            launch_time TEXT,
            raw_export_values_json TEXT NOT NULL,
            first_seen_at TEXT NOT NULL,
            last_seen_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS ranking_rows (
            module TEXT NOT NULL,
            region TEXT NOT NULL,
            snapshot_date TEXT NOT NULL,
            rank INTEGER NOT NULL,
            product_id TEXT NOT NULL,
            source_file TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            PRIMARY KEY(module, region, snapshot_date, rank),
            FOREIGN KEY(product_id) REFERENCES products(product_id),
            FOREIGN KEY(source_file) REFERENCES source_exports(source_file)
        );
        CREATE TABLE IF NOT EXISTS image_assets (
            url TEXT PRIMARY KEY,
            local_path TEXT,
            sha256 TEXT,
            content_type TEXT,
            byte_count INTEGER,
            status TEXT NOT NULL,
            attempts INTEGER NOT NULL DEFAULT 0,
            last_error TEXT,
            updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS ranking_product_idx ON ranking_rows(product_id);
        CREATE INDEX IF NOT EXISTS products_category_idx ON products(region,category);
        """
    )
    return db


def import_search(db: sqlite3.Connection, path: Path, region: str,
                  snapshot_date: str, first_rank: int, last_rank: int) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    expected = last_rank - first_rank + 1
    actual = sheet.max_row - 1
    if actual != expected:
        raise ValueError(f"{path.name}: expected {expected} rows, found {actual}")

    rows = []
    seen_ids = set()
    stamp = now()
    for offset, values in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
        values = list(values)
        if len(values) < 20:
            raise ValueError(f"{path.name}: row {offset + 2} has only {len(values)} columns")
        title, image_url, detail_url = values[0], values[15], values[16]
        match = PRODUCT_ID.search(str(detail_url or ""))
        if not title or not image_url or not match:
            raise ValueError(f"{path.name}: missing title/image/product ID at row {offset + 2}")
        product_id = match.group(1)
        if product_id in seen_ids:
            raise ValueError(f"{path.name}: duplicate product ID {product_id}")
        seen_ids.add(product_id)
        rows.append((first_rank + offset, offset + 2, product_id, values))

    file_name = str(path.resolve())
    file_hash = sha256_file(path)
    old = db.execute("SELECT sha256 FROM source_exports WHERE source_file=?", (file_name,)).fetchone()
    if old and old[0] != file_hash:
        raise ValueError(f"{path.name}: source file changed since import")

    with db:
        db.execute(
            """INSERT OR REPLACE INTO source_exports
            VALUES(?,?,?,?,?,?,?,?,?)""",
            (file_name, file_hash, region, "商品搜索", snapshot_date,
             first_rank, last_rank, actual, stamp),
        )
        for rank, source_row, product_id, v in rows:
            db.execute(
                """INSERT INTO products VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(product_id) DO UPDATE SET
                  title=excluded.title, status=excluded.status,
                  shop_name=excluded.shop_name, region=excluded.region,
                  category=excluded.category, price=excluded.price,
                  commission_rate=excluded.commission_rate,
                  day7_sold=excluded.day7_sold, day7_gmv=excluded.day7_gmv,
                  total_sold=excluded.total_sold, total_gmv=excluded.total_gmv,
                  author_count=excluded.author_count,
                  author_order_rate=excluded.author_order_rate,
                  video_count=excluded.video_count, live_count=excluded.live_count,
                  image_url=excluded.image_url,
                  fastmoss_detail_url=excluded.fastmoss_detail_url,
                  tiktok_detail_url=excluded.tiktok_detail_url,
                  shop_detail_url=excluded.shop_detail_url,
                  launch_time=excluded.launch_time,
                  raw_export_values_json=excluded.raw_export_values_json,
                  last_seen_at=excluded.last_seen_at""",
                (product_id, str(v[0]), str(v[1] or ""), str(v[2] or ""),
                 str(v[3] or region), str(v[4] or ""), str(v[5] or ""),
                 str(v[6] or ""), v[7], v[8], v[9], v[10], v[11],
                 str(v[12] or ""), v[13], v[14], str(v[15]), str(v[16]),
                 str(v[17] or ""), str(v[18] or ""), str(v[19] or ""),
                 json.dumps(v, ensure_ascii=False, default=str), stamp, stamp),
            )
            db.execute(
                """INSERT OR REPLACE INTO ranking_rows
                VALUES(?,?,?,?,?,?,?)""",
                ("商品搜索", region, snapshot_date, rank, product_id,
                 file_name, source_row),
            )
            db.execute(
                """INSERT OR IGNORE INTO image_assets(url,status,updated_at)
                VALUES(?,?,?)""",
                (str(v[15]), "pending", stamp),
            )
    return {"source": file_name, "sha256": file_hash, "rows": actual,
            "first_rank": first_rank, "last_rank": last_rank}


def image_extension(content_type: str, data: bytes) -> str:
    if data.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return ".webp"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return ".gif"
    if (content_type == "image/avif" and data[4:8] == b"ftyp" and
            (data[8:12] in (b"avif", b"avis") or b"avif" in data[12:32])):
        return ".avif"
    if content_type.startswith("image/svg+xml") and b"<svg" in data[:1024]:
        return ".svg"
    raise ValueError(f"unrecognized image format: {content_type}")


def download_one(db: sqlite3.Connection, root: Path, url: str) -> dict:
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme != "https" or not parsed.hostname:
        raise ValueError(f"invalid image URL: {url[:120]}")
    if parsed.hostname not in {
        "s.500fd.com",
        "p16-oec-general.ttcdn-us.com",
        "p19-oec-general-useast5.ttcdn-us.com",
        "p16-oec-general-useast5.ttcdn-us.com",
        "p19-oec-general.ttcdn-us.com",
    }:
        raise ValueError(f"unreviewed image host: {parsed.hostname}")
    last_error = None
    for attempt in range(1, 4):
        try:
            request = urllib.request.Request(
                url, headers={"User-Agent": "Mozilla/5.0", "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8"}
            )
            with urllib.request.urlopen(request, timeout=20) as response:
                content_type = response.headers.get("Content-Type", "").split(";")[0].lower()
                data = response.read(20 * 1024 * 1024 + 1)
            if len(data) > 20 * 1024 * 1024:
                raise ValueError("image exceeds 20 MB limit")
            if not content_type.startswith("image/"):
                raise ValueError(f"unexpected Content-Type: {content_type}")
            ext = image_extension(content_type, data)
            digest = hashlib.sha256(data).hexdigest()
            url_key = hashlib.sha256(url.encode("utf-8")).hexdigest()
            relative = Path("images") / url_key[:2] / f"{url_key}{ext}"
            target = root / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            temporary = target.with_name(target.name + ".part")
            temporary.write_bytes(data)
            os.replace(temporary, target)
            with db:
                db.execute(
                    """UPDATE image_assets SET local_path=?,sha256=?,content_type=?,
                    byte_count=?,status='downloaded',attempts=attempts+?,
                    last_error=NULL,updated_at=? WHERE url=?""",
                    (str(relative).replace("\\", "/"), digest, content_type,
                     len(data), attempt, now(), url),
                )
            return {"url": url, "bytes": len(data), "path": str(target)}
        except (OSError, ValueError, urllib.error.URLError) as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            if attempt < 3:
                time.sleep(attempt * 2)
    with db:
        db.execute(
            """UPDATE image_assets SET status='failed',attempts=attempts+3,
            last_error=?,updated_at=? WHERE url=?""",
            (last_error, now(), url),
        )
    raise RuntimeError(f"image failed after 3 attempts: {url[:120]}: {last_error}")


def download_images(db: sqlite3.Connection, root: Path, limit: int | None,
                    retry_failed: bool = False) -> None:
    statuses = "('pending','failed')" if retry_failed else "('pending')"
    query = f"SELECT url FROM image_assets WHERE status IN {statuses} ORDER BY rowid"
    if limit is not None:
        query += f" LIMIT {int(limit)}"
    urls = [row[0] for row in db.execute(query)]
    downloaded = 0
    failures = 0
    for number, url in enumerate(urls, 1):
        try:
            download_one(db, root, url)
            downloaded += 1
        except (RuntimeError, ValueError) as exc:
            failures += 1
            with db:
                db.execute("""UPDATE image_assets SET status='failed',last_error=?,
                    updated_at=? WHERE url=? AND status!='downloaded'""",
                    (str(exc)[:1000], now(), url))
            print(json.dumps({"image_failure": url, "error": str(exc)},
                             ensure_ascii=True), flush=True)
        if number == 1 or number % 100 == 0 or number == len(urls):
            print(json.dumps({"processed_this_run": number, "downloaded": downloaded,
                              "failed": failures, "target": len(urls)},
                             ensure_ascii=True), flush=True)
        time.sleep(0.2)


def status(db: sqlite3.Connection) -> dict:
    return {
        "sources": db.execute("SELECT COUNT(*) FROM source_exports").fetchone()[0],
        "products": db.execute("SELECT COUNT(*) FROM products").fetchone()[0],
        "ranking_rows": db.execute("SELECT COUNT(*) FROM ranking_rows").fetchone()[0],
        "min_rank": db.execute("SELECT MIN(rank) FROM ranking_rows").fetchone()[0],
        "max_rank": db.execute("SELECT MAX(rank) FROM ranking_rows").fetchone()[0],
        "images": dict(db.execute("SELECT status,COUNT(*) FROM image_assets GROUP BY status")),
        "ranking_gaps": db.execute(
            """SELECT COUNT(*) FROM (
            SELECT rank, LAG(rank) OVER (ORDER BY rank) AS previous
            FROM ranking_rows WHERE module='商品搜索' AND region='US'
            ) WHERE previous IS NOT NULL AND rank != previous+1"""
        ).fetchone()[0],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    sub = parser.add_subparsers(dest="command", required=True)
    ingest = sub.add_parser("import-search")
    ingest.add_argument("--date", required=True)
    ingest.add_argument("--region", default="US")
    ingest.add_argument("--range", action="append", nargs=3,
                        metavar=("FIRST", "LAST", "FILE"), required=True)
    images = sub.add_parser("download-images")
    images.add_argument("--limit", type=int)
    images.add_argument("--retry-failed", action="store_true")
    sub.add_parser("status")
    args = parser.parse_args()
    db = connect(args.root)
    if args.command == "import-search":
        for first, last, file_path in args.range:
            result = import_search(db, Path(file_path), args.region,
                                   args.date, int(first), int(last))
            print(json.dumps(result, ensure_ascii=True), flush=True)
        print(json.dumps(status(db), ensure_ascii=True), flush=True)
    elif args.command == "download-images":
        download_images(db, args.root, args.limit, args.retry_failed)
        print(json.dumps(status(db), ensure_ascii=True), flush=True)
    else:
        print(json.dumps(status(db), ensure_ascii=True), flush=True)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": f"{type(error).__name__}: {error}"},
                         ensure_ascii=True), file=sys.stderr, flush=True)
        raise
