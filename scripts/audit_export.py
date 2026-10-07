"""Read-only checks for a FastMoss product-search XLSX export."""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


PRODUCT_ID = re.compile(r"/e-commerce/detail/(\d+)")


def audit(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    ids: list[str] = []
    missing = Counter()
    first = None
    last = None
    for row in sheet.iter_rows(min_row=2, values_only=True):
        title, image_url, detail_url = row[0], row[15], row[16]
        match = PRODUCT_ID.search(str(detail_url or ""))
        product_id = match.group(1) if match else None
        if not title:
            missing["title"] += 1
        if not image_url:
            missing["image_url"] += 1
        if not product_id:
            missing["product_id"] += 1
        if image_url and not str(image_url).startswith("https://"):
            missing["image_https"] += 1
        if product_id:
            ids.append(product_id)
        if first is None:
            first = product_id
        last = product_id
    counts = Counter(ids)
    return {
        "path": str(path),
        "sheet": sheet.title,
        "data_rows": sheet.max_row - 1,
        "column_count": sheet.max_column,
        "unique_product_ids": len(counts),
        "duplicate_product_ids": sum(n - 1 for n in counts.values() if n > 1),
        "missing_or_invalid": dict(missing),
        "first_product_id": first,
        "last_product_id": last,
        "header_has_replacement_chars": any(
            "\ufffd" in str(cell.value) for cell in next(sheet.iter_rows(max_row=1))
        ),
    }


if __name__ == "__main__":
    print(json.dumps(audit(Path(sys.argv[1])), ensure_ascii=True))
