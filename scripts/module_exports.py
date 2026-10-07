"""Import and reconcile all six FastMoss product-module XLSX exports."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from collections import defaultdict
from datetime import datetime, timezone
import os
from pathlib import Path

from openpyxl import load_workbook
from catalog import connect


ROOT = Path(os.environ.get("FASTMOSS_DATA_ROOT", str(Path(__file__).resolve().parents[1] / "data")))
ID_PATTERN = re.compile(r"/e-commerce/detail/(\d+)")
SOURCES = [
    ("商品搜索", 1, 3000, "US_商品搜索_001-300.xlsx", None),
    ("商品搜索", 3001, 5000, "US_商品搜索_301-500.xlsx", None),
    ("销量榜", 1, 500, "US_销量榜_001-050.xlsx", None),
    ("新品榜", 1, 203, "US_新品榜_001-021.xlsx", "2026-09-27"),
    ("全托管商品榜", 1, 500, "US_全托管商品榜_001-050.xlsx", None),
    ("热推榜", 1, 500, "US_热推榜_001-050.xlsx", None),
    ("视频商品榜", 1, 3000, "US_视频商品榜_001-300.xlsx", None),
    ("视频商品榜", 3001, 5000, "US_视频商品榜_301-500.xlsx", None),
]


def setup(db: sqlite3.Connection) -> None:
    db.executescript("""
        CREATE TABLE IF NOT EXISTS module_sources (
            source_file TEXT PRIMARY KEY, sha256 TEXT NOT NULL,
            module TEXT NOT NULL, region TEXT NOT NULL,
            snapshot_date TEXT NOT NULL, filter_date TEXT,
            first_rank INTEGER NOT NULL, last_rank INTEGER NOT NULL,
            row_count INTEGER NOT NULL, imported_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS module_rows (
            module TEXT NOT NULL, region TEXT NOT NULL,
            snapshot_date TEXT NOT NULL, rank INTEGER NOT NULL,
            product_id TEXT, id_source TEXT NOT NULL,
            title TEXT NOT NULL, image_url TEXT, image_source TEXT NOT NULL,
            detail_url TEXT, source_file TEXT NOT NULL,
            source_row INTEGER NOT NULL, raw_headers_json TEXT NOT NULL,
            raw_values_json TEXT NOT NULL, imported_at TEXT NOT NULL,
            PRIMARY KEY(module, region, snapshot_date, rank),
            FOREIGN KEY(source_file) REFERENCES module_sources(source_file)
        );
        CREATE INDEX IF NOT EXISTS module_rows_product_idx ON module_rows(product_id);
        CREATE INDEX IF NOT EXISTS module_rows_image_idx ON module_rows(image_url);
    """)


def find_column(headers: list[str], options: tuple[str, ...]) -> int | None:
    names = [name.replace(" ", "") for name in headers]
    for option in options:
        if option in names:
            return names.index(option)
    return None


def import_file(db: sqlite3.Connection, path: Path, module: str, region: str,
                snapshot_date: str, first_rank: int, last_rank: int,
                filter_date: str | None) -> dict:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    headers = [str(cell.value or "") for cell in next(sheet.iter_rows(max_row=1))]
    title_col = find_column(headers, ("商品名称", "商品标题", "商品"))
    image_col = find_column(headers, ("商品图片", "商品封面", "商品封面链接"))
    id_col = find_column(headers, ("商品ID",))
    detail_cols = [i for i, name in enumerate(headers)
                   if "FastMoss" in name and "商品详情页" in name]
    detail_col = detail_cols[0] if detail_cols else None
    if title_col is None:
        raise ValueError(f"{path.name}: title column missing")
    expected = last_rank - first_rank + 1
    actual = sheet.max_row - 1
    if actual != expected:
        raise ValueError(f"{path.name}: expected {expected} rows, got {actual}")
    source = str(path.resolve())
    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    old = db.execute("SELECT sha256 FROM module_sources WHERE source_file=?", (source,)).fetchone()
    if old and old[0] != file_hash:
        raise ValueError(f"{path.name}: source changed after import")
    stamp = datetime.now(timezone.utc).isoformat()
    entries = []
    seen_ids = set()
    for offset, values in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
        values = list(values)
        title = str(values[title_col] or "").strip()
        if not title:
            raise ValueError(f"{path.name}: missing title at row {offset + 2}")
        image = str(values[image_col] or "") if image_col is not None else ""
        detail = str(values[detail_col] or "") if detail_col is not None else ""
        explicit = values[id_col] if id_col is not None else None
        match = ID_PATTERN.search(detail)
        product_id = str(explicit) if explicit else (match.group(1) if match else None)
        if explicit and match and str(explicit) != match.group(1):
            raise ValueError(f"{path.name}: ID/detail mismatch at row {offset + 2}")
        if product_id:
            if product_id in seen_ids:
                raise ValueError(f"{path.name}: duplicate product ID {product_id}")
            seen_ids.add(product_id)
        entries.append((first_rank + offset, offset + 2, product_id,
                        title, image, detail, values))
    with db:
        db.execute("INSERT OR REPLACE INTO module_sources VALUES(?,?,?,?,?,?,?,?,?,?)",
                   (source, file_hash, module, region, snapshot_date, filter_date,
                    first_rank, last_rank, actual, stamp))
        for rank, source_row, product_id, title, image, detail, values in entries:
            db.execute("INSERT OR REPLACE INTO module_rows VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                       (module, region, snapshot_date, rank, product_id,
                        "export" if product_id else "unresolved",
                        title, image or None, "export" if image else "unresolved",
                        detail or None, source, source_row,
                        json.dumps(headers, ensure_ascii=False),
                        json.dumps(values, ensure_ascii=False, default=str), stamp))
            if image:
                db.execute("INSERT OR IGNORE INTO image_assets(url,status,updated_at) VALUES(?,'pending',?)",
                           (image, stamp))
    return {"module": module, "file": path.name, "rows": actual,
            "with_id": len(seen_ids),
            "with_image": sum(bool(entry[4]) for entry in entries),
            "sha256": file_hash}


def reconcile(db: sqlite3.Connection, region: str, snapshot_date: str) -> dict:
    image_to_ids: dict[str, set[str]] = defaultdict(set)
    id_to_images: dict[str, set[str]] = defaultdict(set)
    for product_id, image in db.execute(
        "SELECT product_id,image_url FROM module_rows WHERE region=? AND snapshot_date=?",
        (region, snapshot_date)
    ):
        if product_id and image:
            image_to_ids[image].add(product_id)
            id_to_images[product_id].add(image)
    for product_id, image in db.execute("SELECT product_id,image_url FROM products WHERE region=?", (region,)):
        if product_id and image:
            image_to_ids[image].add(product_id)
            id_to_images[product_id].add(image)
    ids_filled = 0
    images_filled = 0
    with db:
        rows = list(db.execute(
            "SELECT module,rank,image_url FROM module_rows "
            "WHERE region=? AND snapshot_date=? AND product_id IS NULL AND image_url IS NOT NULL",
            (region, snapshot_date)
        ))
        for module, rank, image in rows:
            candidates = image_to_ids.get(image, set())
            if len(candidates) == 1:
                db.execute("UPDATE module_rows SET product_id=?,id_source='matched_image' "
                           "WHERE module=? AND region=? AND snapshot_date=? AND rank=?",
                           (next(iter(candidates)), module, region, snapshot_date, rank))
                ids_filled += 1
        rows = list(db.execute(
            "SELECT module,rank,product_id FROM module_rows "
            "WHERE region=? AND snapshot_date=? AND image_url IS NULL AND product_id IS NOT NULL",
            (region, snapshot_date)
        ))
        for module, rank, product_id in rows:
            candidates = id_to_images.get(product_id, set())
            if len(candidates) == 1:
                db.execute("UPDATE module_rows SET image_url=?,image_source='matched_product_id' "
                           "WHERE module=? AND region=? AND snapshot_date=? AND rank=?",
                           (next(iter(candidates)), module, region, snapshot_date, rank))
                images_filled += 1
    return {"ids_filled_from_exact_image_url": ids_filled,
            "images_filled_from_exact_product_id": images_filled}


def status(db: sqlite3.Connection, region: str, snapshot_date: str) -> dict:
    modules = {}
    for module, count, missing_id, missing_image, lo, hi in db.execute(
        "SELECT module,COUNT(*),SUM(product_id IS NULL),SUM(image_url IS NULL),MIN(rank),MAX(rank) "
        "FROM module_rows WHERE region=? AND snapshot_date=? GROUP BY module",
        (region, snapshot_date)
    ):
        modules[module] = {"rows": count, "missing_id": missing_id,
                           "missing_image_url": missing_image,
                           "min_rank": lo, "max_rank": hi}
    return {"sources": db.execute("SELECT COUNT(*) FROM module_sources").fetchone()[0],
            "rows": sum(item["rows"] for item in modules.values()),
            "unique_known_products": db.execute(
                "SELECT COUNT(DISTINCT product_id) FROM module_rows WHERE product_id IS NOT NULL"
            ).fetchone()[0],
            "modules": modules,
            "images": dict(db.execute("SELECT status,COUNT(*) FROM image_assets GROUP BY status"))}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=("import", "status"))
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--date", default="2026-09-30")
    args = parser.parse_args()
    db = connect(args.root)
    setup(db)
    if args.command == "import":
        raw_dir = args.root / "raw" / args.date
        for module, first, last, name, filter_date in SOURCES:
            print(json.dumps(import_file(db, raw_dir / name, module, "US", args.date,
                                         first, last, filter_date), ensure_ascii=True), flush=True)
        print(json.dumps(reconcile(db, "US", args.date), ensure_ascii=True), flush=True)
    print(json.dumps(status(db, "US", args.date), ensure_ascii=True), flush=True)


if __name__ == "__main__":
    main()

